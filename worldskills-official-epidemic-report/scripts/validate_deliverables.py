"""Validate the required WorldSkills epidemic-report deliverable structure."""

from __future__ import annotations

import argparse
import csv
import sys
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


SKILL_DIR = Path(__file__).resolve().parents[1]
LIST_FILE = SKILL_DIR / "references" / "参赛国家和地区名单.csv"
MAIN_HEADERS = [
    "国家/地区", "重点疾病", "近期疫情情况", "高发/暴发/异常上升/新发",
    "病例（包括疑似、确诊）", "死亡数", "最近更新时间", "信息来源", "建议世赛关注程度",
]
EVIDENCE_HEADERS = [
    "来源类型", "疾病", "国家/地区", "日期", "来源文档", "来源 URL", "访问日期", "核实状态",
]
FOCUS_HEADERS = ["国家/地区", "重点疾病", "近期疫情情况", "疫情性质"]
DISEASE_HEADERS = ["疾病", "重点国家", "近期疫情情况", "趋势判断", "与赛事相关性"]


def participant_count() -> int:
    with LIST_FILE.open("r", encoding="utf-8-sig", newline="") as handle:
        return sum(
            1
            for row in csv.DictReader(handle)
            if (row.get("地区（代码 - 英文名称）") or "").strip()
            or (row.get("中文标准名称") or "").strip()
        )


def validate_workbook(path: Path, count: int) -> list[str]:
    errors: list[str] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    for name in ("疫情摸底底表", "网络证据清单"):
        if name not in workbook.sheetnames:
            errors.append(f"Excel缺少工作表：{name}")
    if errors:
        return errors

    main = workbook["疫情摸底底表"]
    evidence = workbook["网络证据清单"]
    if list(next(main.values)) != MAIN_HEADERS:
        errors.append("Excel主表列名或顺序不符合要求")
    if list(next(evidence.values)) != EVIDENCE_HEADERS:
        errors.append("网络证据清单列名或顺序不符合要求")

    source_urls = {str(row[7]).strip() for row in main.iter_rows(min_row=2, values_only=True) if row[7]}
    evidence_urls = {str(row[5]).strip() for row in evidence.iter_rows(min_row=2, values_only=True) if row[5]}
    missing = sorted(source_urls - evidence_urls)
    if missing:
        errors.append(f"主表有{len(missing)}个来源URL未写入网络证据清单")

    if "逐国检索记录" in workbook.sheetnames:
        coverage = workbook["逐国检索记录"].max_row - 1
        if coverage != count:
            errors.append(f"逐国检索记录为{coverage}行，动态名单为{count}行")
    return errors


def validate_document(path: Path, count: int) -> list[str]:
    errors: list[str] = []
    document = Document(path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    required_headings = [
        "一、重点国家", "二、重点疫情摸底", "（一）蚊媒及其他虫媒传染病",
        "（二）呼吸道传染病", "（三）肠道及食源性传染病", "（四）新发少见及高致病性传染病",
    ]
    for heading in required_headings:
        if heading not in text:
            errors.append(f"Word缺少章节：{heading}")
    if f"摸底范围：{count}个" not in text:
        errors.append("Word摸底范围未使用动态名单数")
    headers = [[cell.text for cell in table.rows[0].cells] for table in document.tables if table.rows]
    if not headers or headers[0] != FOCUS_HEADERS:
        errors.append("Word重点国家表列名或顺序不符合要求")
    if len(headers) != 5 or any(header != DISEASE_HEADERS for header in headers[1:]):
        errors.append("Word疾病分类小表数量、列名或顺序不符合要求")
    if document.tables and len(document.tables[0].rows) - 1 > 10:
        errors.append("Word重点国家超过10个")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--docx", type=Path, required=True)
    args = parser.parse_args()
    errors = validate_workbook(args.xlsx, participant_count()) + validate_document(args.docx, participant_count())
    if errors:
        print("\n".join(errors), file=sys.stderr)
        return 1
    print(f"交付校验通过：动态名单{participant_count()}个，Excel与Word结构符合规范。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
