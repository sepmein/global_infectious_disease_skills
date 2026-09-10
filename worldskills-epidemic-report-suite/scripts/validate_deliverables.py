#!/usr/bin/env python3
"""Validate the structural contract of daily or baseline Word/Excel outputs."""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from openpyxl import load_workbook

from validate_scope import (
    DEFAULT_DISEASES,
    DEFAULT_PARTICIPANTS,
    validate_diseases,
    validate_participants,
)


sys.stdout.reconfigure(encoding="utf-8")
sys.stderr.reconfigure(encoding="utf-8")


BASELINE_MAIN_HEADERS = [
    "国家/地区",
    "重点疾病",
    "近期疫情情况",
    "高发/暴发/异常上升/新发",
    "病例（包括疑似、确诊）",
    "死亡数",
    "最近更新时间",
    "信息来源",
    "建议世赛关注程度",
]
NETWORK_EVIDENCE_HEADERS = [
    "来源类型",
    "疾病",
    "国家/地区",
    "日期",
    "来源文档",
    "来源 URL",
    "访问日期",
    "核实状态",
]
FOCUS_HEADERS = ["国家/地区", "重点疾病", "近期疫情情况", "疫情性质"]
DAILY_SHEETS = {"当日变化", "经核验疫情底表", "来源登记", "检索覆盖记录", "重叠与冲突"}
BASELINE_SHEETS = {
    "疫情摸底底表",
    "网络证据清单",
    "来源登记",
    "经核验疫情底表",
    "逐国检索记录",
    "本地文件清单",
    "重叠与冲突",
    "数据字典",
}
PLACEHOLDER_PATTERN = re.compile(r"\*{2,}|待定|待确认|TBD|TODO|PLACEHOLDER", re.IGNORECASE)
DATE_PATTERN = re.compile(r"^\d{4}年(?:[1-9]|1[0-2])月(?:[1-9]|[12]\d|3[01])日$")
DAILY_DATE_PATTERN = re.compile(r"^（\d{4}年(?:[1-9]|1[0-2])月(?:[1-9]|[12]\d|3[01])日）$")
LEVEL1_PATTERN = re.compile(r"^[一二三四五六七八九十]+、")
LEVEL2_PATTERN = re.compile(r"^（[一二三四五六七八九十]+）")
LEVEL3_PATTERN = re.compile(r"^\d+\.\s*")


def font_name(run, paragraph) -> str:
    r_pr = run._r.rPr
    if r_pr is not None and r_pr.rFonts is not None:
        for attribute in (qn("w:eastAsia"), qn("w:ascii"), qn("w:hAnsi")):
            value = r_pr.rFonts.get(attribute)
            if value:
                return value
    if run.font.name:
        return run.font.name
    return paragraph.style.font.name or ""


def font_size(run, paragraph) -> float | None:
    size = run.font.size or paragraph.style.font.size
    return size.pt if size else None


def first_text_run(paragraph):
    return next((run for run in paragraph.runs if run.text.strip()), None)


def normalized_font(value: str) -> str:
    return re.sub(r"[\s_-]+", "", value or "").casefold()


def font_matches(value: str, allowed: set[str]) -> bool:
    normalized = normalized_font(value)
    return any(normalized_font(candidate) in normalized for candidate in allowed)


def paragraph_spacing_ok(paragraph, require_first_line: bool) -> list[str]:
    errors: list[str] = []
    fmt = paragraph.paragraph_format
    before = fmt.space_before.pt if fmt.space_before else 0
    after = fmt.space_after.pt if fmt.space_after else 0
    if abs(before) > 0.2 or abs(after) > 0.2:
        errors.append("段前段后必须为0 pt")
    line = fmt.line_spacing
    line_pt = line.pt if hasattr(line, "pt") else None
    if fmt.line_spacing_rule != WD_LINE_SPACING.EXACTLY or line_pt is None or abs(line_pt - 27) > 0.3:
        errors.append("行距必须为固定值27 pt")
    if require_first_line:
        p_pr = paragraph._p.pPr
        ind = p_pr.ind if p_pr is not None else None
        chars = ind.get(qn("w:firstLineChars")) if ind is not None else None
        indent = fmt.first_line_indent.pt if fmt.first_line_indent else None
        if chars != "200" and (indent is None or abs(indent - 32) > 1.5):
            errors.append("首行缩进必须为2个字符")
        if fmt.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
            errors.append("正文必须两端对齐")
    return errors


def heading_indent_ok(paragraph) -> list[str]:
    fmt = paragraph.paragraph_format
    p_pr = paragraph._p.pPr
    ind = p_pr.ind if p_pr is not None else None
    chars = ind.get(qn("w:firstLineChars")) if ind is not None else None
    indent = fmt.first_line_indent.pt if fmt.first_line_indent else None
    if chars != "200" and (indent is None or abs(indent - 32) > 1.5):
        return ["标题前必须空两格"]
    return []


def validate_paragraph_font(paragraph, allowed_fonts: set[str], size_pt: float, require_bold: bool | None) -> list[str]:
    run = first_text_run(paragraph)
    if run is None:
        return []
    errors: list[str] = []
    actual_font = font_name(run, paragraph)
    if not font_matches(actual_font, allowed_fonts):
        errors.append(f"字体应为{'/'.join(sorted(allowed_fonts))}，实际为{actual_font or '未设置'}")
    actual_size = font_size(run, paragraph)
    if actual_size is None or abs(actual_size - size_pt) > 0.2:
        errors.append(f"字号应为{size_pt:g} pt，实际为{actual_size if actual_size is not None else '未设置'}")
    if require_bold is not None and bool(run.bold) != require_bold:
        errors.append("加粗设置不符合要求")
    return errors


def validate_official_document_format(document, mode: str) -> list[str]:
    if mode == "baseline":
        return validate_baseline_template_format(document)

    errors: list[str] = []
    for section_number, section in enumerate(document.sections, start=1):
        if section.left_margin is None or abs(section.left_margin.mm - 28) > 0.6:
            errors.append(f"Word 第{section_number}节左页边距不是28 mm")
        if section.right_margin is None or abs(section.right_margin.mm - 28) > 0.6:
            errors.append(f"Word 第{section_number}节右页边距不是28 mm")

    title_text = (
        "世界技能大赛相关国家传染病疫情监测日报"
        if mode == "daily"
        else "世界技能大赛参赛国家和地区重点传染病疫情报告"
    )
    paragraphs = document.paragraphs
    title_index = next((index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == title_text), None)
    if title_index is None:
        return errors + ["Word 未找到可校验的报告标题"]
    title = paragraphs[title_index]
    title_errors = validate_paragraph_font(title, {"方正小标宋简体"}, 22, False)
    if title.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.CENTER:
        title_errors.append("标题必须居中")
    if title.paragraph_format.line_spacing_rule not in (None, WD_LINE_SPACING.SINGLE):
        title_errors.append("标题必须为单倍行距")
    line = title.paragraph_format.line_spacing
    if isinstance(line, (int, float)) and abs(float(line) - 1.0) > 0.01:
        title_errors.append("标题必须为单倍行距")
    if mode == "daily" and not title.text.startswith("\u3000\u3000"):
        title_errors.append("日报标题文字前必须保留两个全角空格")
    errors.extend(f"Word 标题：{item}" for item in title_errors)

    next_nonempty = next((index for index in range(title_index + 1, len(paragraphs)) if paragraphs[index].text.strip()), None)
    if next_nonempty is None:
        return errors + ["Word 标题后缺少报告日期和正文" if mode == "daily" else "Word 标题后缺少主送单位和正文"]

    issuer_index: int | None = None
    date_indices: list[int] = []
    if mode == "daily":
        date_index = next_nonempty
        date_paragraph = paragraphs[date_index]
        if date_index != title_index + 1:
            errors.append("Word 日报日期必须紧接标题单独成行")
        if not DAILY_DATE_PATTERN.fullmatch(date_paragraph.text.strip()):
            errors.append("Word 日报日期必须使用不补零的括号格式，如（2026年9月7日）")
        if date_paragraph.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.CENTER:
            errors.append("Word 日报日期必须居中")
        for item in validate_paragraph_font(date_paragraph, {"楷体_GB2312", "楷体GB", "楷体"}, 16, True):
            errors.append(f"Word 日报日期：{item}")
        for item in paragraph_spacing_ok(date_paragraph, require_first_line=False):
            errors.append(f"Word 日报日期：{item}")
        body_start = date_index + 1
        body_end = len(paragraphs)
    else:
        if next_nonempty == title_index + 1:
            errors.append("Word 标题与主送单位之间必须空一行")
        addressee = paragraphs[next_nonempty]
        addressee_text = addressee.text.strip()
        if not addressee_text.endswith("："):
            errors.append("Word 主送单位必须使用全称并以全角冒号结束")
        if PLACEHOLDER_PATTERN.search(addressee_text):
            errors.append("Word 主送单位包含占位文字")
        if addressee.paragraph_format.alignment not in (None, WD_ALIGN_PARAGRAPH.LEFT):
            errors.append("Word 主送单位必须居左")
        if addressee.paragraph_format.left_indent and abs(addressee.paragraph_format.left_indent.pt) > 0.2:
            errors.append("Word 主送单位必须顶格")
        for item in validate_paragraph_font(addressee, {"仿宋_GB2312", "仿宋GB", "仿宋"}, 16, None):
            errors.append(f"Word 主送单位：{item}")
        for item in paragraph_spacing_ok(addressee, require_first_line=False):
            errors.append(f"Word 主送单位：{item}")

        date_indices = [index for index, paragraph in enumerate(paragraphs) if DATE_PATTERN.fullmatch(paragraph.text.strip())]
        if not date_indices:
            errors.append("Word 缺少不补零的中文成文日期，如2026年9月7日")
            date_index = len(paragraphs)
        else:
            date_index = date_indices[-1]
            date_paragraph = paragraphs[date_index]
            p_pr = date_paragraph._p.pPr
            ind = p_pr.ind if p_pr is not None else None
            right_chars = ind.get(qn("w:rightChars")) if ind is not None else None
            right_indent = date_paragraph.paragraph_format.right_indent.pt if date_paragraph.paragraph_format.right_indent else None
            if right_chars != "400" and (right_indent is None or not 60 <= right_indent <= 68):
                errors.append("Word 成文日期必须右空四字编排")
            for item in validate_paragraph_font(date_paragraph, {"仿宋_GB2312", "仿宋GB", "仿宋"}, 16, None):
                errors.append(f"Word 成文日期：{item}")
            for item in paragraph_spacing_ok(date_paragraph, require_first_line=False):
                errors.append(f"Word 成文日期：{item}")

            issuer_index = next((index for index in range(date_index - 1, next_nonempty, -1) if paragraphs[index].text.strip()), None)
            if issuer_index is None:
                errors.append("Word 成文日期前缺少发文单位落款")
            else:
                issuer = paragraphs[issuer_index]
                if PLACEHOLDER_PATTERN.search(issuer.text.strip()):
                    errors.append("Word 发文单位包含占位文字")
                for item in validate_paragraph_font(issuer, {"仿宋_GB2312", "仿宋GB", "仿宋"}, 16, None):
                    errors.append(f"Word 发文单位：{item}")
        body_start = next_nonempty + 1
        body_end = min(date_index, len(paragraphs))

    for index in range(body_start, body_end):
        paragraph = paragraphs[index]
        text = paragraph.text.strip()
        if not text or text.startswith("附件：") or index == issuer_index:
            continue
        if LEVEL1_PATTERN.match(text) or text.startswith("附："):
            paragraph_errors = validate_paragraph_font(paragraph, {"黑体"}, 16, False)
            paragraph_errors += paragraph_spacing_ok(paragraph, require_first_line=False)
            paragraph_errors += heading_indent_ok(paragraph)
            errors.extend(f"Word 一级标题“{text[:20]}”：{item}" for item in paragraph_errors)
        elif LEVEL2_PATTERN.match(text):
            paragraph_errors = validate_paragraph_font(paragraph, {"楷体_GB2312", "楷体GB", "楷体"}, 16, True)
            paragraph_errors += paragraph_spacing_ok(paragraph, require_first_line=False)
            paragraph_errors += heading_indent_ok(paragraph)
            errors.extend(f"Word 二级标题“{text[:20]}”：{item}" for item in paragraph_errors)
        elif LEVEL3_PATTERN.match(text):
            paragraph_errors = validate_paragraph_font(paragraph, {"仿宋_GB2312", "仿宋GB", "仿宋"}, 16, None)
            paragraph_errors += paragraph_spacing_ok(paragraph, require_first_line=False)
            paragraph_errors += heading_indent_ok(paragraph)
            errors.extend(f"Word 三级标题“{text[:20]}”：{item}" for item in paragraph_errors)
        else:
            paragraph_errors = validate_paragraph_font(paragraph, {"仿宋_GB2312", "仿宋GB", "仿宋"}, 16, None)
            paragraph_errors += paragraph_spacing_ok(paragraph, require_first_line=True)
            errors.extend(f"Word 正文段“{text[:20]}”：{item}" for item in paragraph_errors)

    if mode == "daily":
        source_index = next((index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "附：信息来源"), None)
        if source_index is None:
            errors.append("Word 日报缺少“附：信息来源”")
        else:
            source = paragraphs[source_index]
            has_page_break = bool(source.paragraph_format.page_break_before)
            if source_index > 0:
                has_page_break = has_page_break or any(
                    element.get(qn("w:type")) == "page"
                    for element in paragraphs[source_index - 1]._p.iter(qn("w:br"))
                )
            if not has_page_break:
                errors.append("Word 日报“附：信息来源”必须另起一页")

        review_tags = {
            "w:commentRangeStart": "批注锚点",
            "w:commentReference": "批注引用",
            "w:ins": "插入修订",
            "w:del": "删除修订",
            "w:moveFrom": "移动修订",
            "w:moveTo": "移动修订",
            "w:rPrChange": "字符格式修订",
            "w:pPrChange": "段落格式修订",
        }
        for tag, label in review_tags.items():
            if any(True for _ in document._element.iter(qn(tag))):
                errors.append(f"Word 日报不得保留{label}")
        if any(
            element.get(qn("w:val"), "1") not in {"0", "false", "off", "none"}
            for element in document._element.iter(qn("w:strike"))
        ):
            errors.append("Word 日报不得保留删除线")
        if any(
            element.get(qn("w:val"), "none") != "none"
            for element in document._element.iter(qn("w:highlight"))
        ):
            errors.append("Word 日报不得保留文字高亮")

    attachment_indices = [index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip().startswith("附件：")]
    for attachment_index in attachment_indices:
        if attachment_index == 0 or paragraphs[attachment_index - 1].text.strip():
            errors.append("Word 附件说明与正文之间必须空一行")
        if date_indices:
            issuer_index = next((index for index in range(date_indices[-1] - 1, attachment_index, -1) if paragraphs[index].text.strip()), None)
            if issuer_index is not None:
                blank_count = sum(1 for index in range(attachment_index + 1, issuer_index) if not paragraphs[index].text.strip())
                if blank_count < 3:
                    errors.append("Word 附件说明与落款之间必须至少空三行")

    for table_number, table in enumerate(document.tables, start=1):
        if table.rows:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            if tr_pr.find(qn("w:tblHeader")) is None:
                errors.append(f"Word 第{table_number}个表格首行未设置为重复表头")
        for row_number, row in enumerate(table.rows, start=1):
            for cell in row.cells:
                paragraph = next((p for p in cell.paragraphs if p.text.strip()), None)
                if paragraph is None:
                    continue
                allowed = {"黑体"} if row_number == 1 else {"仿宋_GB2312", "仿宋GB", "仿宋"}
                require_bold = False if row_number == 1 else None
                for item in validate_paragraph_font(paragraph, allowed, 16, require_bold):
                    errors.append(f"Word 第{table_number}个表格第{row_number}行：{item}")
    return errors


def validate_baseline_template_format(document) -> list[str]:
    """Validate the retained baseline template's stable layout contract."""
    errors: list[str] = []
    for section_number, section in enumerate(document.sections, start=1):
        if section.page_width is None or abs(section.page_width.mm - 210) > 0.8:
            errors.append(f"Word 第{section_number}节页面宽度不是A4")
        if section.page_height is None or abs(section.page_height.mm - 297) > 0.8:
            errors.append(f"Word 第{section_number}节页面高度不是A4")
        for label, actual, expected in (
            ("左", section.left_margin, 28),
            ("右", section.right_margin, 28),
            ("上", section.top_margin, 25.4),
            ("下", section.bottom_margin, 25.4),
        ):
            if actual is None or abs(actual.mm - expected) > 0.8:
                errors.append(f"Word 第{section_number}节{label}页边距未沿用摸底报告模板")

    paragraphs = document.paragraphs
    title_text = "世界技能大赛参赛国家和地区重点传染病疫情报告"
    title = next((p for p in paragraphs if p.text.strip() == title_text), None)
    if title is None:
        return errors + ["Word 未找到摸底报告模板标题"]
    title_errors = validate_paragraph_font(title, {"方正小标宋简体"}, 22, False)
    if title.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.CENTER:
        title_errors.append("标题必须居中")
    if title.paragraph_format.line_spacing_rule not in (None, WD_LINE_SPACING.SINGLE):
        title_errors.append("标题必须为单倍行距")
    errors.extend(f"Word 标题：{item}" for item in title_errors)

    role_specs = {
        "一、重点国家": ({"黑体"}, False),
        "二、重点疫情摸底": ({"黑体"}, False),
        "（一）蚊媒及其他虫媒传染病": ({"楷体_GB2312", "楷体GB", "楷体"}, True),
        "（二）肠道及食源性传染病": ({"楷体_GB2312", "楷体GB", "楷体"}, True),
        "（三）呼吸道传染病": ({"楷体_GB2312", "楷体GB", "楷体"}, True),
        "（四）新发少见及高致病性传染病": ({"楷体_GB2312", "楷体GB", "楷体"}, True),
    }
    by_text = {p.text.strip(): p for p in paragraphs if p.text.strip()}
    for text, (fonts, bold) in role_specs.items():
        paragraph = by_text.get(text)
        if paragraph is None:
            continue
        paragraph_errors = validate_paragraph_font(paragraph, fonts, 16, bold)
        paragraph_errors += paragraph_spacing_ok(paragraph, require_first_line=False)
        errors.extend(f"Word 模板标题“{text}”：{item}" for item in paragraph_errors)

    for text in ("一、疾病概述", "二、全球流行概况", "三、重点国家流行情况", "四、风险因素"):
        matches = [p for p in paragraphs if p.text.strip() == text]
        for paragraph in matches:
            paragraph_errors = validate_paragraph_font(paragraph, {"黑体"}, 16, False)
            paragraph_errors += paragraph_spacing_ok(paragraph, require_first_line=False)
            errors.extend(f"Word 疾病分节“{text}”：{item}" for item in paragraph_errors)
    return errors


def row_values(sheet, row_number: int = 1) -> list[str]:
    return [str(cell.value or "").strip() for cell in sheet[row_number]]


def validate_coverage_sheet(sheet, expected_count: int) -> list[str]:
    errors: list[str] = []
    headers = row_values(sheet)
    code_index = next(
        (headers.index(name) for name in ("country_code", "地区代码", "国家/地区代码") if name in headers),
        None,
    )
    if code_index is None:
        return [f"{sheet.title} 缺少 country_code/地区代码列"]
    codes = {
        str(row[code_index].value or "").strip()
        for row in sheet.iter_rows(min_row=2)
        if str(row[code_index].value or "").strip()
    }
    if len(codes) != expected_count:
        errors.append(f"{sheet.title} 有 {len(codes)} 个唯一地区代码，动态名单为 {expected_count} 个")
    return errors


def validate_document(path: Path, mode: str, participant_count: int, require_official_format: bool = False) -> list[str]:
    errors: list[str] = []
    document = Document(path)
    text = "\n".join(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
    table_headers = [
        [cell.text.strip() for cell in table.rows[0].cells]
        for table in document.tables
        if table.rows
    ]
    if mode == "daily":
        required = [
            "世界技能大赛相关国家传染病疫情监测日报",
            "检索说明",
            "一、近期赛事相关国家和地区疫情总体概况",
            "二、当日新增及变化疫情动态",
            "三、其他未更新传染病疫情",
            "四、对上海市传染病输入风险综合研判",
            "附：信息来源",
        ]
        for item in required:
            if item not in text:
                errors.append(f"Word 日报缺少：{item}")
    else:
        required = [
            "世界技能大赛参赛国家和地区重点传染病疫情报告",
            "一、重点国家",
            "二、重点疫情摸底",
            "（一）蚊媒及其他虫媒传染病",
            "（二）肠道及食源性传染病",
            "（三）呼吸道传染病",
            "（四）新发少见及高致病性传染病",
        ]
        for item in required:
            if item not in text:
                errors.append(f"Word 摸底报告缺少：{item}")
        if "摸底范围" not in text or str(participant_count) not in text:
            errors.append("Word 摸底范围未使用动态名单数量")
        if not table_headers or table_headers[0] != FOCUS_HEADERS:
            errors.append("Word 重点国家表表头或顺序不符合要求")
        elif len(document.tables[0].rows) - 1 > 10:
            errors.append("Word 重点国家超过 10 个")
        for item in ("一、疾病概述", "二、全球流行概况", "三、重点国家流行情况", "四、风险因素"):
            if item not in text:
                errors.append(f"Word 摸底报告缺少逐病种结构：{item}")
        forbidden_template_text = (
            "具体的病种名称",
            "某参赛国家",
            "填写该份监测报告的真实链接",
            "所有数字和情境均为虚构",
            "占位示例",
        )
        for item in forbidden_template_text:
            if item in text:
                errors.append(f"Word 摸底报告残留模板说明或占位文字：{item}")
    if require_official_format:
        errors.extend(validate_official_document_format(document, mode))
    return errors


def validate_workbook(
    path: Path,
    mode: str,
    participant_count: int,
    disease_records: list[dict[str, str]] | None = None,
) -> list[str]:
    errors: list[str] = []
    workbook = load_workbook(path, read_only=True, data_only=False)
    required_sheets = DAILY_SHEETS if mode == "daily" else BASELINE_SHEETS
    missing = sorted(required_sheets - set(workbook.sheetnames))
    if missing:
        errors.append(f"Excel 缺少工作表：{'、'.join(missing)}")
        return errors
    coverage_name = "检索覆盖记录" if mode == "daily" else "逐国检索记录"
    errors.extend(validate_coverage_sheet(workbook[coverage_name], participant_count))
    if mode == "baseline":
        if row_values(workbook["疫情摸底底表"]) != BASELINE_MAIN_HEADERS:
            errors.append("疫情摸底底表列名或顺序不符合要求")
        if row_values(workbook["网络证据清单"]) != NETWORK_EVIDENCE_HEADERS:
            errors.append("网络证据清单列名或顺序不符合要求")
        allowed_diseases = {row["disease"] for row in (disease_records or [])}
        for sheet_name, header_name in (("疫情摸底底表", "重点疾病"), ("网络证据清单", "疾病")):
            sheet = workbook[sheet_name]
            headers = row_values(sheet)
            if header_name not in headers:
                continue
            disease_index = headers.index(header_name)
            invalid = sorted({
                str(row[disease_index].value or "").strip()
                for row in sheet.iter_rows(min_row=2)
                if str(row[disease_index].value or "").strip()
                and str(row[disease_index].value or "").strip() not in allowed_diseases
            })
            if invalid:
                errors.append(f"{sheet_name} 包含重点病种表外疾病：{'、'.join(invalid)}")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    if cell.value in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}:
                        errors.append(f"{sheet.title}!{cell.coordinate} 包含公式错误：{cell.value}")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--mode", choices=("daily", "baseline"), required=True)
    parser.add_argument("--participants", type=Path, default=DEFAULT_PARTICIPANTS)
    parser.add_argument("--diseases", type=Path, default=DEFAULT_DISEASES)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument(
        "--require-official-format",
        action="store_true",
        help="Validate the required Word layout: daily official format or baseline retained template.",
    )
    args = parser.parse_args()
    try:
        participant_count = len(validate_participants(args.participants))
        disease_records = validate_diseases(args.diseases) if args.mode == "baseline" else None
        errors = validate_document(args.docx, args.mode, participant_count, args.require_official_format)
        errors.extend(validate_workbook(args.xlsx, args.mode, participant_count, disease_records))
    except Exception as exc:
        print(f"无法校验交付件：{type(exc).__name__}: {exc}", file=sys.stderr)
        return 1
    if errors:
        print("\n".join(errors), file=sys.stderr)
        return 1
    print(f"交付结构校验通过：mode={args.mode}，动态名单={participant_count}。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
