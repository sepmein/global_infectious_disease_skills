#!/usr/bin/env python3
"""Run smoke tests for the skill's deterministic validation scripts."""

from __future__ import annotations

import csv
import json
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from openpyxl import Workbook


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent


def set_font(run, name: str, size: float, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_fixed_spacing(paragraph, first_line: bool = False) -> None:
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(27)
    if first_line:
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_pr = paragraph._p.get_or_add_pPr()
        ind = p_pr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            p_pr.append(ind)
        ind.set(qn("w:firstLineChars"), "200")


def add_formatted_paragraph(document, text: str, font: str, size: float = 16, bold: bool | None = None, first_line: bool = False):
    paragraph = document.add_paragraph()
    set_fixed_spacing(paragraph, first_line=first_line)
    run = paragraph.add_run(text)
    set_font(run, font, size, bold)
    return paragraph


def create_official_daily_fixture(docx_path: Path, xlsx_path: Path, participants: list[dict[str, str]]) -> None:
    document = Document()
    for section in document.sections:
        section.left_margin = Mm(28)
        section.right_margin = Mm(28)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    title.paragraph_format.line_spacing = 1.0
    run = title.add_run("\u3000\u3000世界技能大赛相关国家传染病疫情监测日报")
    set_font(run, "方正小标宋简体", 22, False)

    report_date = add_formatted_paragraph(document, "（2026年9月7日）", "楷体_GB2312", bold=True)
    report_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formatted_paragraph(
        document,
        "检索说明：本期覆盖动态名单，并与上一版日报和历史底表比较。",
        "仿宋_GB2312",
        first_line=True,
    )
    for heading in (
        "一、近期赛事相关国家和地区疫情总体概况",
        "二、当日新增及变化疫情动态",
        "三、其他未更新传染病疫情",
        "四、对上海市传染病输入风险综合研判",
        "附：信息来源",
    ):
        heading_paragraph = add_formatted_paragraph(document, heading, "黑体", bold=False, first_line=True)
        if heading == "附：信息来源":
            heading_paragraph.paragraph_format.page_break_before = True
        if heading.startswith("一、"):
            add_formatted_paragraph(document, "（一）二级标题测试", "楷体_GB2312", bold=True, first_line=True)
            add_formatted_paragraph(document, "1. 三级标题测试", "仿宋_GB2312", bold=False, first_line=True)
        add_formatted_paragraph(document, "本段用于验证正文的字体、缩进和固定行距。", "仿宋_GB2312", first_line=True)
    document.save(docx_path)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in ("当日变化", "经核验疫情底表", "来源登记", "检索覆盖记录", "重叠与冲突"):
        sheet = workbook.create_sheet(name)
        if name == "检索覆盖记录":
            sheet.append(["country_code", "coverage_status"])
            for participant in participants:
                sheet.append([participant["country_code"], "CHECKED_NO_SIGNAL"])
        else:
            sheet.append(["id"])
    workbook.save(xlsx_path)


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text(
        "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows),
        encoding="utf-8",
    )


def run(case_id: str, command: list[str], expected_code: int = 0) -> None:
    completed = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
    if completed.returncode != expected_code:
        raise AssertionError(
            f"{case_id} failed: expected {expected_code}, got {completed.returncode}\n"
            f"stdout={completed.stdout}\nstderr={completed.stderr}"
        )
    print(f"{case_id} PASS")


def main() -> int:
    python = sys.executable
    run("BVT_001_SCOPE_DEFAULT", [python, str(SCRIPT_DIR / "validate_scope.py")])

    from validate_deliverables import validate_baseline_template_format

    template = SKILL_DIR / "assets" / "世界技能大赛传染病疫情摸底报告_写作模板.docx"
    template_errors = validate_baseline_template_format(Document(template))
    if template_errors:
        raise AssertionError(f"BVT_011 baseline template format errors: {template_errors}")
    print("BVT_011_BASELINE_TEMPLATE_FORMAT PASS")

    intro_index = SKILL_DIR / "references" / "disease-intro-index.csv"
    with intro_index.open("r", encoding="utf-8-sig", newline="") as handle:
        intro_rows = list(csv.DictReader(handle))
    missing_intro_files = [
        row["重点病种"]
        for row in intro_rows
        if not (SKILL_DIR / "references" / "疾病简介" / row["疾病简介文件"]).is_file()
    ]
    if len(intro_rows) != 22 or missing_intro_files:
        raise AssertionError(
            f"BVT_012 intro mapping count={len(intro_rows)} missing={missing_intro_files}"
        )
    print("BVT_012_DISEASE_INTRO_INDEX PASS")

    with tempfile.TemporaryDirectory(prefix="epidemic-skill-test-") as temp_name:
        temp = Path(temp_name)
        participants = temp / "participants.csv"
        with participants.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["地区（代码 - 英文名称）", "中文标准名称"])
            writer.writerow(["AA - Alpha", "甲国"])
            writer.writerow(["BB - Beta", "乙国"])

        coverage_ok = temp / "coverage-ok.jsonl"
        write_jsonl(
            coverage_ok,
            [
                {
                    "country_code": "AA",
                    "coverage_status": "VERIFIED_EVENT",
                    "checked_at": "2026-09-07T08:00:00+08:00",
                    "channels_checked": ["LOCAL_WHO", "NATIONAL_OFFICIAL"],
                    "event_ids": ["E0001"],
                },
                {
                    "country_code": "BB",
                    "coverage_status": "CHECKED_NO_SIGNAL",
                    "checked_at": "2026-09-07T08:00:00+08:00",
                    "channels_checked": ["WHO_REGIONAL", "NATIONAL_OFFICIAL"],
                },
            ],
        )
        run(
            "BVT_002_COVERAGE_VALID",
            [python, str(SCRIPT_DIR / "validate_coverage.py"), str(coverage_ok), "--participants", str(participants)],
        )

        coverage_bad = temp / "coverage-bad.jsonl"
        write_jsonl(
            coverage_bad,
            [
                {"country_code": "AA", "coverage_status": "NOT_CHECKED"},
                {
                    "country_code": "BB",
                    "coverage_status": "CHECKED_NO_SIGNAL",
                    "checked_at": "2026-09-07T08:00:00+08:00",
                    "channels_checked": ["WHO_REGIONAL"],
                },
            ],
        )
        run(
            "BVT_003_COVERAGE_REJECTS_NOT_CHECKED",
            [python, str(SCRIPT_DIR / "validate_coverage.py"), str(coverage_bad), "--participants", str(participants)],
            expected_code=1,
        )

        sources = temp / "sources.jsonl"
        events = temp / "events.jsonl"
        previous = temp / "previous.jsonl"
        current = temp / "current.jsonl"
        changes = temp / "changes.jsonl"
        write_jsonl(
            sources,
            [
                {
                    "source_id": "S0001",
                    "source_channel": "WEB",
                    "publisher": "Example Ministry of Health",
                    "title": "Outbreak update",
                    "locator": "Table 1",
                    "evidence_grade": "A",
                    "original_url": "https://example.gov/outbreak",
                }
            ],
        )
        write_jsonl(
            events,
            [
                {
                    "event_id": "E0001",
                    "country_code": "AA",
                    "disease_canonical": "示例病",
                    "source_ids": ["S0001"],
                    "inclusion_status": "INCLUDED",
                }
            ],
        )
        write_jsonl(
            previous,
            [
                {
                    "observation_id": "O0001",
                    "event_id": "E0001",
                    "observed_at": "2026-09-06T08:00:00+08:00",
                    "case_count": 10,
                    "death_count": 1,
                    "case_metric": "cumulative",
                    "case_definition": "confirmed",
                    "affected_areas": "甲地",
                    "source_ids": ["S0001"],
                }
            ],
        )
        write_jsonl(
            current,
            [
                {
                    "observation_id": "O0002",
                    "event_id": "E0001",
                    "observed_at": "2026-09-07T08:00:00+08:00",
                    "case_count": 14,
                    "death_count": 1,
                    "case_metric": "cumulative",
                    "case_definition": "confirmed",
                    "affected_areas": "甲地；乙地",
                    "source_ids": ["S0001"],
                }
            ],
        )
        run(
            "BVT_004_EVIDENCE_LINKAGE",
            [
                python,
                str(SCRIPT_DIR / "validate_evidence_package.py"),
                "--sources",
                str(sources),
                "--events",
                str(events),
                "--observations",
                str(current),
            ],
        )
        run(
            "BVT_005_DAILY_CHANGE_DETECTION",
            [
                python,
                str(SCRIPT_DIR / "detect_daily_changes.py"),
                "--current",
                str(current),
                "--previous",
                str(previous),
                "--out",
                str(changes),
            ],
        )
        change_types = {row["change_type"] for row in read_jsonl(changes)}
        if change_types != {"CASE_INCREASE", "AREA_EXPANSION"}:
            raise AssertionError(f"BVT_005 unexpected change types: {sorted(change_types)}")
        print("BVT_006_CHANGE_TYPES PASS")

        official_docx = temp / "official-daily.docx"
        official_xlsx = temp / "official-daily.xlsx"
        create_official_daily_fixture(
            official_docx,
            official_xlsx,
            [
                {"country_code": "AA", "country_name_en": "Alpha", "country_name_zh": "甲国"},
                {"country_code": "BB", "country_name_en": "Beta", "country_name_zh": "乙国"},
            ],
        )
        run(
            "BVT_007_OFFICIAL_FORMAT_VALID",
            [
                python,
                str(SCRIPT_DIR / "validate_deliverables.py"),
                "--mode",
                "daily",
                "--participants",
                str(participants),
                "--docx",
                str(official_docx),
                "--xlsx",
                str(official_xlsx),
                "--require-official-format",
            ],
        )

        invalid_document = Document(official_docx)
        invalid_document.sections[0].left_margin = Mm(20)
        invalid_docx = temp / "invalid-official-daily.docx"
        invalid_document.save(invalid_docx)
        run(
            "BVT_008_OFFICIAL_FORMAT_REJECTS_BAD_MARGIN",
            [
                python,
                str(SCRIPT_DIR / "validate_deliverables.py"),
                "--mode",
                "daily",
                "--participants",
                str(participants),
                "--docx",
                str(invalid_docx),
                "--xlsx",
                str(official_xlsx),
                "--require-official-format",
            ],
            expected_code=1,
        )

        bold_document = Document(official_docx)
        bold_heading = next(paragraph for paragraph in bold_document.paragraphs if paragraph.text.startswith("一、"))
        bold_heading.runs[0].bold = True
        bold_docx = temp / "invalid-bold-heading.docx"
        bold_document.save(bold_docx)
        run(
            "BVT_009_OFFICIAL_FORMAT_REJECTS_BOLD_HEITI",
            [
                python,
                str(SCRIPT_DIR / "validate_deliverables.py"),
                "--mode",
                "daily",
                "--participants",
                str(participants),
                "--docx",
                str(bold_docx),
                "--xlsx",
                str(official_xlsx),
                "--require-official-format",
            ],
            expected_code=1,
        )

        indent_document = Document(official_docx)
        indent_heading = next(paragraph for paragraph in indent_document.paragraphs if paragraph.text.startswith("二、"))
        indent = indent_heading._p.get_or_add_pPr().find(qn("w:ind"))
        if indent is not None:
            indent.set(qn("w:firstLineChars"), "0")
        indent_docx = temp / "invalid-heading-indent.docx"
        indent_document.save(indent_docx)
        run(
            "BVT_010_OFFICIAL_FORMAT_REJECTS_HEADING_WITHOUT_INDENT",
            [
                python,
                str(SCRIPT_DIR / "validate_deliverables.py"),
                "--mode",
                "daily",
                "--participants",
                str(participants),
                "--docx",
                str(indent_docx),
                "--xlsx",
                str(official_xlsx),
                "--require-official-format",
            ],
            expected_code=1,
        )
    return 0


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


if __name__ == "__main__":
    raise SystemExit(main())
